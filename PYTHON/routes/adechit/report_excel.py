"""Excel output of the same report blocks used by the print and PDF renderers."""
import re
from decimal import Decimal
from html import escape
from io import BytesIO
import unicodedata

from .domain import DomainError, require

MIMETYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def column_key(column, index):
    if column.get('key'):
        return column['key']
    text = unicodedata.normalize('NFKD', column['label'].replace('-\n', ''))
    text = ''.join(char for char in text if not unicodedata.combining(char))
    return re.sub(r'[^A-Z0-9]+', '_', text.upper()).strip('_') or ('NR_CRT' if index == 0 else f'COLOANA_{index + 1}')


def configured_widths(repo, model):
    """Read shared settings on each export; no DDL, writes or schema synchronization."""
    if repo.sqlite or not model.get('report_name'):
        return {}
    from mysql.connector.errors import ProgrammingError
    try:
        rows = repo.query('SELECT CHEIE_COLOANA, LATIME FROM AVACONT_COMUN.RAPOARTE_EXCEL_COLOANE WHERE NUME_RAPORT=%s',
                          (model['report_name'],))
    except ProgrammingError as error:
        if error.errno == 1146:  # The user has not created the settings table yet.
            return {}
        raise
    return {row['CHEIE_COLOANA']: float(row['LATIME']) for row in rows if 0 < float(row['LATIME']) <= 255}


def table_schemas(model):
    """Repeated tables (groups, children and totals) share one column selection."""
    schemas = []
    for block in model['blocks']:
        if block['t'] == 'table':
            labels = tuple(column['label'] for column in block['cols'])
            if labels not in schemas:
                schemas.append(labels)
    return schemas


def column_options(model):
    return [(f'{table}:{index}', label.replace('-\n', '').replace('\n', ' ') or ('Nr. crt.' if index == 0 else f'Coloana {index + 1}'))
            for table, labels in enumerate(table_schemas(model)) for index, label in enumerate(labels)]


def dialog_html(model):
    options = column_options(model)
    choices = ''.join(f'<label><input type="checkbox" name="column" value="{key}" checked><span>{escape(label)}</span></label>'
                      for key, label in options)
    return ('<dialog id="report-excel-dialog" class="ade-export" aria-labelledby="report-excel-title" aria-describedby="report-excel-note">'
            '<form id="report-excel-form"><header><h2 id="report-excel-title">Export Excel</h2>'
            '<button id="report-excel-close" type="button" aria-label="Închide">✕</button></header>'
            '<div class="ade-export-content"><p id="report-excel-note">Alegeți coloanele pentru export. Antetul se exportă întotdeauna.</p>'
            + ('<label class="ade-export-all"><input id="report-excel-all" type="checkbox" checked><span>Toate coloanele</span></label>'
               f'<div class="ade-export-columns">{choices}</div>' if options else '<p>Acest raport nu are coloane de selectat. Se exportă integral.</p>')
            + '<p id="report-excel-error" role="alert" hidden></p></div><footer>'
            '<button id="report-excel-cancel" type="button">Renunță</button>'
            '<button id="report-excel-save" type="submit">Exportă Excel</button></footer></form></dialog>')


def selected_columns(model, args):
    allowed = {key for key, _ in column_options(model)}
    raw = args.get('columns')
    if raw is None:
        return allowed
    chosen = set(raw.split(',')) if raw else set()
    require(chosen <= allowed, 'EXPORT_COLUMNS', 'Selecția coloanelor este invalidă.', 400)
    require(bool(chosen) or not allowed, 'EXPORT_COLUMNS', 'Selectați cel puțin o coloană.', 400)
    return chosen


def xlsx_bytes(model, args, column_widths=None):
    chosen = selected_columns(model, args)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise DomainError('EXCEL_DEPENDENCY', 'Exportul Excel necesită instalarea dependențelor din requirements-adechit.txt pe server.', 503) from error

    book = Workbook()
    sheet = book.active
    sheet.title = 'Raport'
    book.properties.title = model['title']
    schemas = table_schemas(model)
    width = max(2, max((sum(f'{table}:{index}' in chosen for index in range(len(labels)))
                        for table, labels in enumerate(schemas)), default=0))
    line = 1
    configured_dimensions = {}
    edge = Side(style='thin', color='A6A6A6')
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    tints = {'d': 'F9EDED', 'c': 'E6EDD7', 'z': 'DFE5ED'}

    def write(row, column, value, span=1, bold=False, align='left', fill=None, grid=False, numeric=False):
        text = str(value if value is not None else '')
        cell = sheet.cell(row, column)
        # Identifiers and user text must stay literal, including text starting with '='.
        cell.value = text
        cell.data_type = 's'
        cell.number_format = '@'
        if numeric and re.fullmatch(r'-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?', text):
            cell.value = Decimal(text.replace('.', '').replace(',', '.'))
            cell.number_format = '#,##0.00' if ',' in text else '#,##0'
        cell.font = Font(name='Calibri', size=11, bold=bold)
        cell.alignment = Alignment(horizontal=align, vertical='top', wrap_text=True)
        if fill:
            cell.fill = PatternFill('solid', fgColor=fill)
        if grid:
            cell.border = border
        if span > 1:
            sheet.merge_cells(start_row=row, start_column=column, end_row=row, end_column=column + span - 1)

    def full_row(text, bold=False, align='left'):
        nonlocal line
        write(line, 1, text, width, bold=bold, align=align)
        sheet.row_dimensions[line].height = max(18, 16 * len(str(text).split('\n')))
        line += 1

    for block in model['blocks']:
        kind = block['t']
        if kind == 'table':
            columns = block['cols']
            table = schemas.index(tuple(column['label'] for column in columns))
            kept = [index for index in range(len(columns)) if f'{table}:{index}' in chosen]
            if not kept:
                continue
            if block.get('head', True):
                for output, index in enumerate(kept, 1):
                    column = columns[index]
                    configured = (column_widths or {}).get(column_key(column, index))
                    if configured is not None:
                        configured_dimensions[output] = max(configured_dimensions.get(output, 0), configured)
                    write(line, output, column['label'], bold=True, align='center', grid=True,
                          fill=tints.get(column.get('tint'), 'ECECEC'))
                    sheet.column_dimensions[get_column_letter(output)].width = max(
                        sheet.column_dimensions[get_column_letter(output)].width, min(55, max(14, column['w'] / 100)))
                sheet.row_dimensions[line].height = 16 * max(len(columns[index]['label'].split('\n')) for index in kept)
                if sheet.freeze_panes is None:
                    sheet.freeze_panes = f'A{line + 1}'
                line += 1
            for row in block['rows']:
                source = 0
                for cell in row['cells']:
                    text, span = (cell['t'], int(cell.get('span', 1))) if isinstance(cell, dict) else (cell, 1)
                    overlap = [index for index in kept if source <= index < source + span]
                    if overlap:
                        column = columns[source]
                        total = row.get('cls') in ('total', 'grand', 'sub')
                        write(line, kept.index(overlap[0]) + 1, text, len(overlap), bold=total, grid=True,
                              align={'r': 'right', 'c': 'center'}.get(column.get('align'), 'left'),
                              fill='ECECEC' if total else tints.get(column.get('tint')),
                              numeric=column.get('align') == 'r')
                    source += span
                line += 1
            line += 1
        elif kind == 'letterhead':
            for text in block['left']:
                full_row(text, bold=True)
            if block.get('right'):
                full_row(block['right'], bold=True, align='right')
            line += 1
        elif kind == 'title':
            full_row(block['text'], bold=True, align='center')
            for text in block.get('sub', []):
                full_row(text, align='center')
        elif kind == 'text':
            full_row(block['text'], bold=block.get('bold', False), align=block.get('align', 'left'))
        elif kind == 'box':
            if block.get('title'):
                full_row(block['title'], bold=True)
            for label, value in block['rows']:
                write(line, 1, label, bold=True)
                write(line, 2, value, width - 1)
                line += 1
            line += 1
        elif kind == 'sign':
            full_row('    |    '.join(block['heads']))
            line += 2
        elif kind in ('spacer', 'break'):
            line += 1
        else:
            raise DomainError('REPORT_BLOCK', 'Bloc de raport necunoscut.', 500)
    if model.get('footer'):
        full_row(model['footer'])
    for index in range(1, width + 1):
        dimension = sheet.column_dimensions[get_column_letter(index)]
        dimension.width = configured_dimensions.get(index, max(dimension.width, 20))
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = 'landscape' if model['landscape'] else 'portrait'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
